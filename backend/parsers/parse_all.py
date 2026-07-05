#!/usr/bin/env python3
"""Parse all three faculty timetables into a single seed JSON file."""

import json, re, sys, os
from datetime import datetime

FOBE_XLSX = os.path.expanduser("~/Downloads/mids semester2  2026 fbe.xlsx")
ART_PDF_TXT = None  # We'll use pdftotext output
EDUCATION_DOCX = os.path.expanduser("~/Downloads/SECOND SEMESTER MIDSEM TIMETABLE.docx")

SESSION_TIMES = {
    1: ("08:15", "09:15"),
    2: ("10:00", "11:00"),
    3: ("11:45", "12:45"),
    4: ("13:30", "14:30"),
    5: ("15:15", "16:15"),
    6: ("17:00", "18:00"),
}

DAY_DATES = {
    "monday": "2026-07-06",
    "tuesday": "2026-07-07",
    "wednesday": "2026-07-08",
    "thursday": "2026-07-09",
    "friday": "2026-07-10",
}

def parse_session_number(text):
    text = text.lower().strip()
    for pat, num in [
        (r'session\s*1', 1), (r'8[:.]15', 1),
        (r'session\s*2', 2), (r'10[:.]00', 2),
        (r'session\s*3', 3), (r'11[:.]45', 3),
        (r'session\s*4', 4), (r'1[:.]30', 4), (r'13[:.]30', 4),
        (r'session\s*5', 5), (r'3[:.]15', 5), (r'15[:.]15', 5),
        (r'session\s*6', 6), (r'5[:.]00', 6), (r'17[:.]00', 6),
    ]:
        if re.search(pat, text):
            return num
    return None

def parse_day(text):
    text = text.lower()
    for day in DAY_DATES:
        if day in text:
            return day
    return None

def parse_fobe():
    """Parse FOBE Excel file (with invigilators column)."""
    import openpyxl
    wb = openpyxl.load_workbook(FOBE_XLSX)
    ws = wb.active

    exams = []
    venues_set = set()
    staff_set = set()
    current_day = None
    current_session = None

    rows = list(ws.iter_rows(min_row=7, values_only=True))

    for row in rows:
        if not row or len(row) < 7:
            continue

        date_session = str(row[0] or "").strip()
        course_raw = str(row[1] or "").strip()
        examiner = str(row[2] or "").strip()
        year_raw = str(row[3] or "").strip()
        venue_raw = str(row[4] or "").strip()
        index_nums = str(row[5] or "").strip()
        print_count = row[6]
        invigilators_raw = str(row[7] or "").strip() if len(row) > 7 else ""

        if date_session:
            day = parse_day(date_session)
            if day:
                current_day = day
            sess = parse_session_number(date_session)
            if sess:
                current_session = sess

        if not course_raw or course_raw == "None" or course_raw == "COURSE NO/TITLE":
            continue

        course_match = re.match(r'^([A-Z]{2,5}\s*\d{3}[A-Z]?)\s+(.+)$', course_raw)
        if not course_match:
            continue

        course_code = course_match.group(1).strip()
        course_name = course_match.group(2).strip()

        if venue_raw and venue_raw != "None":
            venues_set.add(venue_raw)

        invigilator_names = []
        if invigilators_raw and invigilators_raw != "None":
            parts = re.split(r'[,;]', invigilators_raw)
            for p in parts:
                name = p.strip()
                if name and len(name) > 2:
                    invigilator_names.append(name)
                    staff_set.add(name)

        if examiner and examiner != "None":
            staff_set.add(examiner)

        students = 0
        if print_count and print_count != "None":
            try:
                students = int(float(print_count))
            except (ValueError, TypeError):
                pass

        exam_type = "written"
        index_upper = index_nums.upper()
        if "BYOD" in index_upper:
            exam_type = "BYOD"
        elif "CBE" in index_upper:
            exam_type = "CBE"

        exam = {
            "faculty": "FOBE",
            "course_code": course_code,
            "course_name": course_name,
            "examiner": examiner if examiner != "None" else "",
            "year_group": year_raw if year_raw != "None" else "",
            "venue": venue_raw if venue_raw != "None" else "",
            "day": current_day,
            "date": DAY_DATES.get(current_day, ""),
            "session": current_session,
            "session_time": SESSION_TIMES.get(current_session, ("", "")),
            "students": students,
            "invigilators": invigilator_names,
            "exam_type": exam_type,
        }
        exams.append(exam)

    return exams, list(venues_set), list(staff_set)


def parse_art_pdf():
    """Parse Art Faculty timetable from PDF text."""
    import subprocess
    art_pdf = os.path.expanduser("~/Downloads/SECOND SEMESTER MID-SEMESTER EXAMINATIONS TIMETABLE 2025 (3).pdf")
    result = subprocess.run(["pdftotext", "-layout", art_pdf, "-"], capture_output=True, text=True)
    lines = result.stdout.split("\n")

    exams = []
    venues_set = set()
    staff_set = set()
    current_day = None
    current_venue = None

    days_in_order = ["monday", "tuesday", "wednesday", "thursday", "friday"]
    day_idx = 0

    i = 0
    while i < len(lines):
        line = lines[i].strip()

        if re.match(r'^(MONDAY|TUESDAY|WEDNESDAY|THURSDAY|FRIDAY)\s*$', line, re.IGNORECASE):
            current_day = line.lower()
            i += 1
            continue

        venue_match = re.match(r'^(GF\d?|FF\d?|GFW|FFW)\s+', lines[i]) if lines[i].strip() else None
        if venue_match:
            current_venue = venue_match.group(1).strip()
            venues_set.add(current_venue)

        course_matches = re.finditer(r'([A-Z]{2,5}\s+\d{3}(?:/\d{3})?)', line)
        for m in course_matches:
            code = m.group(1)
            col_pos = m.start()
            session_num = None
            if col_pos < 25:
                session_num = 1
            elif col_pos < 50:
                session_num = 2
            elif col_pos < 75:
                session_num = 3
            elif col_pos < 100:
                session_num = 4
            elif col_pos < 130:
                session_num = 5
            else:
                session_num = 6

            name_after = line[m.end():].strip()
            name_part = re.split(r'\[|\(|$', name_after)[0].strip()

            students = 0
            count_match = re.search(r'\[(\d+)\]|\((\d+)\)', line[m.start():])
            if count_match:
                students = int(count_match.group(1) or count_match.group(2))

            exam = {
                "faculty": "Art",
                "course_code": code.strip(),
                "course_name": name_part[:80] if name_part else "",
                "examiner": "",
                "year_group": "",
                "venue": current_venue or "",
                "day": current_day,
                "date": DAY_DATES.get(current_day, ""),
                "session": session_num,
                "session_time": SESSION_TIMES.get(session_num, ("", "")),
                "students": students,
                "invigilators": [],
                "exam_type": "written",
            }
            exams.append(exam)

        i += 1

    return exams, list(venues_set), list(staff_set)


def parse_education_docx():
    """Parse Education Faculty timetable from DOCX tables."""
    from docx import Document
    doc = Document(EDUCATION_DOCX)

    exams = []
    venues_set = set()
    staff_set = set()

    session_col_map = {3: 1, 4: 2, 6: 3, 8: 4, 9: 5}

    current_day = None

    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if len(cells) < 4:
                continue

            day_text = cells[0].replace("\n", " ").strip()
            day_match = parse_day(day_text)
            if day_match:
                current_day = day_match

            year_raw = cells[1].strip() if len(cells) > 1 else ""
            year_num = ""
            y_match = re.search(r'\d', year_raw)
            if y_match:
                year_num = y_match.group()

            for col_idx in range(3, min(len(cells), 10)):
                cell_text = cells[col_idx].strip()
                if not cell_text or cell_text in ("B\n" * 8 + "R", "E\n" * 8 + "A\n" * 7 + "K"):
                    continue

                session_num = None
                if col_idx == 3:
                    session_num = 1
                elif col_idx == 4:
                    session_num = 2
                elif col_idx == 6:
                    session_num = 3
                elif col_idx == 8:
                    session_num = 4
                elif col_idx == 9:
                    session_num = 5

                if not session_num:
                    continue

                course_match = re.search(r'([A-Z]{2,5}\s*\d{3}(?:/[A-Z]*\s*\d{3})?)', cell_text)
                if not course_match:
                    continue

                code = course_match.group(1)
                remaining = cell_text[course_match.end():].strip()
                remaining = remaining.lstrip(":").strip()

                lines = remaining.split("\n")
                name_parts = []
                venue = ""
                lecturer = ""
                students = 0

                for l in lines:
                    l = l.strip()
                    if not l:
                        continue
                    count_m = re.search(r'\((\d+)\)', l)
                    if count_m:
                        students = int(count_m.group(1))
                        l = l[:count_m.start()].strip()
                    if re.match(r'^(PUB|LRM|NCB|COS)', l):
                        venue = l
                        venues_set.add(venue)
                    elif re.match(r'^(Prof\.|Dr\.|Mr\.|Mrs\.|Miss|William|Benjamin)', l):
                        lecturer = l
                        staff_set.add(lecturer)
                    elif l and not name_parts:
                        name_parts.append(l)
                    elif l:
                        name_parts.append(l)

                course_name = " ".join(name_parts)[:100]

                if "Postgraduate" in day_text or "POSTGRADUATE" in cells[0]:
                    continue

                edu_exam_type = "written"
                cell_upper = cell_text.upper()
                if "ONLINE" in cell_upper:
                    edu_exam_type = "CBE"
                elif "COMPUTER" in cell_upper:
                    edu_exam_type = "CBE"

                exam = {
                    "faculty": "Education",
                    "course_code": code.strip(),
                    "course_name": course_name,
                    "examiner": lecturer,
                    "year_group": f"Year {year_num}" if year_num else "",
                    "venue": venue,
                    "day": current_day,
                    "date": DAY_DATES.get(current_day, ""),
                    "session": session_num,
                    "session_time": SESSION_TIMES.get(session_num, ("", "")),
                    "students": students,
                    "invigilators": [],
                    "exam_type": edu_exam_type,
                }
                exams.append(exam)

    return exams, list(venues_set), list(staff_set)


def deduplicate_exams(exams):
    """Remove duplicate exams (same course+day+session)."""
    seen = set()
    unique = []
    for e in exams:
        key = (e["course_code"], e["day"], e["session"])
        if key not in seen:
            seen.add(key)
            unique.append(e)
        else:
            for existing in unique:
                ekey = (existing["course_code"], existing["day"], existing["session"])
                if ekey == key:
                    if e["venue"] and e["venue"] not in existing["venue"]:
                        existing["venue"] += ", " + e["venue"]
                    for inv in e["invigilators"]:
                        if inv not in existing["invigilators"]:
                            existing["invigilators"].append(inv)
                    if e["students"] > existing["students"]:
                        existing["students"] = e["students"]
                    break
    return unique


def main():
    all_exams = []
    all_venues = set()
    all_staff = set()

    print("Parsing FOBE Excel...", file=sys.stderr)
    fobe_exams, fobe_venues, fobe_staff = parse_fobe()
    print(f"  Found {len(fobe_exams)} exams, {len(fobe_venues)} venues, {len(fobe_staff)} staff", file=sys.stderr)
    all_exams.extend(fobe_exams)
    all_venues.update(fobe_venues)
    all_staff.update(fobe_staff)

    print("Parsing Art PDF...", file=sys.stderr)
    art_exams, art_venues, art_staff = parse_art_pdf()
    print(f"  Found {len(art_exams)} exams, {len(art_venues)} venues, {len(art_staff)} staff", file=sys.stderr)
    all_exams.extend(art_exams)
    all_venues.update(art_venues)
    all_staff.update(art_staff)

    print("Parsing Education DOCX...", file=sys.stderr)
    edu_exams, edu_venues, edu_staff = parse_education_docx()
    print(f"  Found {len(edu_exams)} exams, {len(edu_venues)} venues, {len(edu_staff)} staff", file=sys.stderr)
    all_exams.extend(edu_exams)
    all_venues.update(edu_venues)
    all_staff.update(edu_staff)

    all_exams = deduplicate_exams(all_exams)
    print(f"\nTotal: {len(all_exams)} unique exams, {len(all_venues)} venues, {len(all_staff)} staff", file=sys.stderr)

    seed = {
        "faculties": ["FOBE", "Art", "Education"],
        "exam_period": {
            "name": "Second Semester Mid-Semester 2025/2026",
            "start_date": "2026-07-06",
            "end_date": "2026-07-10",
        },
        "session_times": {str(k): v for k, v in SESSION_TIMES.items()},
        "venues": sorted(all_venues),
        "staff": sorted(all_staff),
        "exams": all_exams,
    }

    json.dump(seed, sys.stdout, indent=2, default=str)


if __name__ == "__main__":
    main()
